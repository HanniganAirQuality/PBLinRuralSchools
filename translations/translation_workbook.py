#!/usr/bin/env python3
"""Shared helpers for converting locale JSON catalogs to and from XLSX."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import tempfile
from pathlib import Path, PurePosixPath
from typing import Any

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as error:  # pragma: no cover - depends on local setup
    if error.name == "openpyxl":
        raise SystemExit(
            "Missing dependency: openpyxl. Install it with "
            "'python -m pip install -r translations/requirements.txt'."
        ) from error
    raise


METADATA_SHEET = "_catalogs"
METADATA_MARKER = "translation-workbook-v1"
INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")


def _load_catalog(path: Path) -> dict[str, str]:
    """Load a flat, string-to-string translation catalog."""
    with path.open(encoding="utf-8-sig") as stream:
        value = json.load(stream)

    if not isinstance(value, dict):
        raise ValueError(f"{path}: the JSON root must be an object")

    catalog: dict[str, str] = {}
    for key, text in value.items():
        if not isinstance(key, str) or not isinstance(text, str):
            raise ValueError(f"{path}: entry {key!r} must have a string value")
        catalog[key] = text
    return catalog


def _json_files(directory: Path) -> dict[PurePosixPath, Path]:
    return {
        PurePosixPath(path.relative_to(directory).as_posix()): path
        for path in sorted(directory.rglob("*.json"))
        if path.is_file()
    }


def _unique_sheet_title(relative_path: PurePosixPath, used: set[str]) -> str:
    base = str(relative_path.with_suffix(""))
    base = INVALID_SHEET_CHARACTERS.sub(" - ", base).strip(" '") or "catalog"
    candidate = base[:31]
    if candidate.casefold() not in used:
        used.add(candidate.casefold())
        return candidate

    digest = hashlib.sha1(str(relative_path).encode("utf-8")).hexdigest()[:7]
    candidate = f"{base[:23]}-{digest}"
    counter = 2
    while candidate.casefold() in used:
        suffix = f"-{counter}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def _write_text_cell(sheet: Worksheet, row: int, column: int, value: str) -> None:
    """Write text without allowing values beginning with '=' to become formulas."""
    cell = sheet.cell(row=row, column=column, value=value)
    cell.data_type = "s"


def json_to_workbook(locales_dir: Path, english_locale: str, output: Path) -> None:
    locales_dir = locales_dir.resolve()
    english_dir = locales_dir / english_locale
    if not english_dir.is_dir():
        raise ValueError(f"English locale directory not found: {english_dir}")

    locale_dirs = sorted(
        (path for path in locales_dir.iterdir() if path.is_dir()),
        key=lambda path: path.name.casefold(),
    )
    locale_dirs = [english_dir] + [path for path in locale_dirs if path != english_dir]
    locale_names = [path.name for path in locale_dirs]
    english_files = _json_files(english_dir)
    if not english_files:
        raise ValueError(f"No JSON files found under {english_dir}")

    workbook = Workbook()
    workbook.remove(workbook.active)
    metadata = workbook.create_sheet(METADATA_SHEET)
    metadata.sheet_state = "hidden"
    metadata.append([METADATA_MARKER, "sheet", "relative_json_path"])
    used_titles = {METADATA_SHEET.casefold()}

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for relative_path, english_path in english_files.items():
        catalogs: dict[str, dict[str, str]] = {
            english_locale: _load_catalog(english_path)
        }
        english_keys = set(catalogs[english_locale])

        for locale_dir in locale_dirs[1:]:
            locale_path = locale_dir / Path(*relative_path.parts)
            catalog = _load_catalog(locale_path) if locale_path.is_file() else {}
            unexpected = set(catalog) - english_keys
            if unexpected:
                sample = ", ".join(repr(key) for key in sorted(unexpected)[:5])
                remainder = len(unexpected) - 5
                more = f" (and {remainder} more)" if remainder > 0 else ""
                print(
                    f"WARNING: {locale_path}: ignoring key(s) absent from the "
                    f"English catalog: {sample}{more}",
                    file=sys.stderr,
                )
            catalogs[locale_dir.name] = catalog

        title = _unique_sheet_title(relative_path, used_titles)
        sheet = workbook.create_sheet(title)
        metadata.append([None, title, relative_path.as_posix()])

        headers = ["Key", *locale_names]
        for column, header in enumerate(headers, start=1):
            _write_text_cell(sheet, 1, column, header)
            sheet.cell(1, column).fill = header_fill
            sheet.cell(1, column).font = header_font

        for row, key in enumerate(catalogs[english_locale], start=2):
            _write_text_cell(sheet, row, 1, key)
            for column, locale_name in enumerate(locale_names, start=2):
                if key in catalogs[locale_name]:
                    _write_text_cell(sheet, row, column, catalogs[locale_name][key])

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions["A"].width = 48
        for column in range(2, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 55
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    # The metadata sheet is first but hidden; open the workbook on the first catalog.
    workbook.active = 1
    output = output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def _catalog_sheet_paths(workbook: Any) -> list[tuple[Worksheet, PurePosixPath]]:
    if METADATA_SHEET not in workbook.sheetnames:
        raise ValueError(
            f"Workbook is missing the hidden {METADATA_SHEET!r} metadata sheet; "
            "create it with json_to_excel.py"
        )

    metadata = workbook[METADATA_SHEET]
    if metadata.cell(1, 1).value != METADATA_MARKER:
        raise ValueError("Workbook metadata format is not recognized")

    result: list[tuple[Worksheet, PurePosixPath]] = []
    for row in metadata.iter_rows(min_row=2, values_only=True):
        _, sheet_name, relative_name = row[:3]
        if not isinstance(sheet_name, str) or not isinstance(relative_name, str):
            continue
        relative_path = PurePosixPath(relative_name)
        if (
            relative_path.is_absolute()
            or ".." in relative_path.parts
            or relative_path.suffix.lower() != ".json"
        ):
            raise ValueError(f"Unsafe catalog path in workbook metadata: {relative_name!r}")
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Catalog sheet listed in metadata is missing: {sheet_name}")
        result.append((workbook[sheet_name], relative_path))
    if not result:
        raise ValueError("Workbook metadata does not list any catalog sheets")
    return result


def _read_headers(sheet: Worksheet) -> list[str]:
    headers: list[str] = []
    for cell in sheet[1]:
        if cell.value is None:
            break
        if not isinstance(cell.value, str) or not cell.value.strip():
            raise ValueError(f"Sheet {sheet.title!r}: headers must be non-empty text")
        headers.append(cell.value.strip())
    if len(headers) < 2 or headers[0] != "Key":
        raise ValueError(
            f"Sheet {sheet.title!r}: first row must start with 'Key' and a locale folder name"
        )
    if len(set(headers[1:])) != len(headers[1:]):
        raise ValueError(f"Sheet {sheet.title!r}: locale headers must be unique")
    for locale in headers[1:]:
        locale_path = Path(locale)
        if (
            locale in {".", ".."}
            or locale_path.is_absolute()
            or len(locale_path.parts) != 1
            or "/" in locale
            or "\\" in locale
        ):
            raise ValueError(
                f"Sheet {sheet.title!r}: unsafe locale folder name {locale!r}"
            )
    return headers


def _cell_text(sheet: Worksheet, row: int, column: int, label: str) -> str | None:
    cell = sheet.cell(row=row, column=column)
    if cell.data_type == "f":
        raise ValueError(
            f"Sheet {sheet.title!r}, row {row}, {label}: formulas are not supported"
        )
    value = cell.value
    if value is None:
        return None
    if not isinstance(value, str):
        raise ValueError(
            f"Sheet {sheet.title!r}, row {row}, {label}: value must be text"
        )
    return value


def _write_json(path: Path, catalog: dict[str, str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    content = json.dumps(catalog, ensure_ascii=False, indent=2) + "\n"
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", newline="\n", dir=path.parent, delete=False
    ) as stream:
        stream.write(content)
        temporary = Path(stream.name)
    temporary.replace(path)


def workbook_to_json(workbook_path: Path, output_dir: Path, english_locale: str) -> None:
    workbook = load_workbook(workbook_path, data_only=False)
    sheets = _catalog_sheet_paths(workbook)
    expected_locales: list[str] | None = None

    for sheet, relative_path in sheets:
        headers = _read_headers(sheet)
        locales = headers[1:]
        if english_locale not in locales:
            raise ValueError(
                f"Sheet {sheet.title!r}: English locale column {english_locale!r} is missing"
            )
        if expected_locales is None:
            expected_locales = locales
        elif locales != expected_locales:
            raise ValueError(
                f"Sheet {sheet.title!r}: locale columns do not match the other sheets"
            )

        catalogs: dict[str, dict[str, str]] = {locale: {} for locale in locales}
        seen_keys: set[str] = set()
        for row in range(2, sheet.max_row + 1):
            key = _cell_text(sheet, row, 1, "Key")
            values = [
                _cell_text(sheet, row, column, locale)
                for column, locale in enumerate(locales, start=2)
            ]
            if key is None and all(value is None for value in values):
                continue
            if key is None or not key:
                raise ValueError(f"Sheet {sheet.title!r}, row {row}: Key is blank")
            if key in seen_keys:
                raise ValueError(f"Sheet {sheet.title!r}, row {row}: duplicate Key {key!r}")
            seen_keys.add(key)

            for locale, value in zip(locales, values):
                if value is not None:
                    catalogs[locale][key] = value

            if key not in catalogs[english_locale]:
                raise ValueError(
                    f"Sheet {sheet.title!r}, row {row}: English value is blank"
                )

        for locale, catalog in catalogs.items():
            destination = output_dir / locale / Path(*relative_path.parts)
            _write_json(destination, catalog)


def format_error(error: Exception) -> str:
    return f"ERROR: {error}"
